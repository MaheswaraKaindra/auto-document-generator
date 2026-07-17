"""Test routes_document.py (Peran 3) — endpoint /documents/sdd, /uat, /generate.

Proses plantuml.jar dan pemanggilan LLM (Claude) selalu di-mock supaya test
tidak bergantung pada Java/jar terpasang, koneksi internet, API key, atau kuota.
"""

import json
import re
from pathlib import Path
from unittest.mock import Mock, patch

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api import routes_document
from app.api.schemas_document import DocumentMetadata
from app.domain.exceptions import ContextWindowExceededError
from app.main import app
from app.services import job_store

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "dummy_data"

# PNG putih ukuran sungguhan (bukan 1x1): _image_attr kini menghitung ukuran
# tampil dari piksel, dan gambar 1 piksel menghasilkan width=0.00in yang tidak
# mewakili diagram nyata mana pun.
def _white_png(width: int = 1600, height: int = 1200) -> bytes:
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buffer, format="PNG")
    return buffer.getvalue()


_MINIMAL_PNG = _white_png()


def _load_fixture(name: str) -> dict:
    return json.loads((FIXTURES_DIR / name).read_text(encoding="utf-8"))


@pytest.fixture(autouse=True)
def isolated_job_db(tmp_path, monkeypatch):
    """Arahkan SQLite ke tmp_path, jangan sentuh data/jobs.db milik pengembang."""
    monkeypatch.setattr(job_store, "DB_PATH", tmp_path / "jobs.db")
    job_store.init_db()


@pytest.fixture
def client():
    return TestClient(app)


def _generate(client, **body) -> dict:
    """POST /documents/generate lalu ambil status job-nya.

    TestClient menjalankan BackgroundTasks SESUDAH response terkirim tapi SEBELUM
    client.post() balik — jadi begitu baris ini selesai, job-nya sudah rampung.
    Itu yang bikin test async ini tetap deterministik tanpa sleep/polling.
    """
    response = client.post("/documents/generate", json={"repositories": [], **body})
    assert response.status_code == 202, response.text
    job_id = response.json()["job_id"]
    return client.get(f"/documents/jobs/{job_id}").json()


@pytest.fixture
def mock_plantuml_ok():
    with patch(
        "app.services.compiler_service._run_plantuml", return_value=_MINIMAL_PNG
    ) as mocked:
        yield mocked


def test_generate_sdd_from_content_returns_docx(client, mock_plantuml_ok):
    data = _load_fixture("document_content_sdd.json")

    response = client.post("/documents/sdd", json=data)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def test_generate_uat_from_content_returns_docx(client, mock_plantuml_ok):
    data = _load_fixture("document_content_uat.json")

    response = client.post("/documents/uat", json=data)

    assert response.status_code == 200


def test_generate_sdd_missing_field_returns_422(client):
    """DocumentContent divalidasi Pydantic -- body tidak lengkap harus ditolak,
    bukan diproses setengah jalan."""
    response = client.post("/documents/sdd", json={"document_type": "SDD"})

    assert response.status_code == 422


def test_generate_sdd_returns_502_when_diagram_render_fails(client):
    from app.domain.exceptions import DiagramRenderError

    data = _load_fixture("document_content_sdd.json")

    with patch(
        "app.services.compiler_service._run_plantuml",
        side_effect=DiagramRenderError("boom"),
    ):
        response = client.post("/documents/sdd", json=data)

    assert response.status_code == 502


def test_generate_full_pipeline_rejects_invalid_document_type(client):
    response = client.post(
        "/documents/generate",
        json={"document_type": "INVALID", "repositories": []},
    )

    assert response.status_code == 422


def test_generate_returns_202_immediately_without_doing_the_work(client, mock_plantuml_ok):
    """Inti dari async: POST balik SEBELUM pipeline jalan.

    Versi lama menahan seluruh pipeline (terukur 191 detik pada repo nyata) di
    satu request HTTP, sementara proxy umumnya memutus di 30-60 detik.
    """
    content = _load_fixture("document_content_sdd.json")

    with patch.object(
        routes_document._llm_service, "generate_document_content", return_value=content
    ):
        response = client.post(
            "/documents/generate", json={"document_type": "SDD", "repositories": []}
        )

    assert response.status_code == 202
    body = response.json()
    assert body["status"] == "queued"
    assert body["status_url"] == f"/documents/jobs/{body['job_id']}"


def test_llm_failure_lands_on_the_job_not_the_request(client):
    """Kegagalan pipeline tidak lagi bisa dikembalikan sebagai status HTTP request
    — requestnya sudah balik lama. Harus mendarat di job."""
    with patch.object(
        routes_document._llm_service,
        "generate_document_content",
        side_effect=RuntimeError("LLM gagal"),
    ):
        job = _generate(client, document_type="SDD")

    assert job["status"] == "failed"
    assert job["error_status"] == 502
    assert "coba lagi" in job["error"].lower()  # sebab tak dikenal: mengulang memang masuk akal


def test_context_window_failure_keeps_its_413_after_going_async(client):
    """Penjaga regresi paling penting dari perpindahan ke async: pemetaan
    exception -> kode HTTP tidak boleh HILANG hanya karena kegagalannya sekarang
    terjadi di latar belakang. Repo kebesaran itu PERMANEN — pengguna harus tetap
    bisa membedakannya dari 502 'coba lagi'. Kalau semua kegagalan job dilaporkan
    sama, kita balik ke penyamaran yang sudah tiga kali diperbaiki."""
    with patch.object(
        routes_document._llm_service,
        "generate_document_content",
        side_effect=ContextWindowExceededError(
            "Repo ini terlalu besar untuk model claude-sonnet-5: metadata kodenya "
            "1,224,476 token, sementara model ini cuma memuat 1,000,000."
        ),
    ):
        job = _generate(client, document_type="SDD")

    assert job["status"] == "failed"
    assert job["error_status"] == 413
    assert "1,224,476" in job["error"]  # sebab aslinya diteruskan, bukan disamarkan
    assert "coba lagi" not in job["error"].lower()


def _docx_text_from_response(response, tmp_path) -> str:
    """Baca docx yang benar-benar dikirim ke klien. Sel tabel ikut diambil --
    hampir semua metadata mendarat di tabel, dan doc.paragraphs melewatkannya."""
    path = tmp_path / "response.docx"
    path.write_bytes(response.content)
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_form_metadata_survives_the_whole_async_round_trip(client, mock_plantuml_ok, tmp_path):
    """Jalur yang benar-benar dipakai end user, sekarang tiga langkah: isian form
    harus menembus POST -> background task -> compiler -> template -> DB -> lalu
    keluar utuh di docx yang diunduh lewat /jobs/{id}/download."""
    content = _load_fixture("document_content_sdd.json")

    with patch.object(
        routes_document._llm_service, "generate_document_content", return_value=content
    ):
        job = _generate(
            client,
            document_type="SDD",
            document_metadata={
                "rfc_number": "RFC-2026-088",
                "business_requestor": "Divisi Operasional",
            },
        )

    assert job["status"] == "done"
    download = client.get(job["download_url"])
    assert download.status_code == 200
    text = _docx_text_from_response(download, tmp_path)
    assert "RFC-2026-088" in text
    assert "Divisi Operasional" in text


def test_every_meta_field_in_templates_exists_in_schema():
    """Penjaga untuk sisi buruk desain fail-open di _MetadataDict.__missing__:
    field yang tidak dikenal TIDAK error, cuma diam-diam keluar sebagai
    *(diisi manual)*. Artinya salah ketik di template (mis. meta.rfc_numbr)
    tidak akan pernah ketahuan saat runtime -- sel-nya cuma kosong selamanya.
    Test ini yang menangkapnya."""
    known_fields = set(DocumentMetadata.model_fields)
    templates_dir = Path(__file__).resolve().parent.parent / "app" / "templates"

    for template in ("sdd_template.md", "uat_template.md"):
        used = set(re.findall(r"meta\.([a-z_]+)", (templates_dir / template).read_text(encoding="utf-8")))
        unknown = used - known_fields
        assert not unknown, f"{template} memakai field yang tidak ada di DocumentMetadata: {sorted(unknown)}"


def test_generate_works_without_metadata(client, mock_plantuml_ok, tmp_path):
    """document_metadata opsional -- request tanpa field ini harus tetap jalan
    dan menghasilkan dokumen berpenanda seperti sebelumnya."""
    content = _load_fixture("document_content_sdd.json")

    with patch.object(
        routes_document._llm_service, "generate_document_content", return_value=content
    ):
        job = _generate(client, document_type="SDD")

    assert job["status"] == "done"
    download = client.get(job["download_url"])
    assert "(diisi manual)" in _docx_text_from_response(download, tmp_path)


def test_download_before_finished_says_not_ready_not_not_found(client):
    """409, bukan 404: job-nya ADA, cuma belum siap. 404 bikin klien mengira
    job_id-nya salah lalu berhenti polling."""
    job_id = job_store.create_job(document_type="SDD", project_name=None)

    response = client.get(f"/documents/jobs/{job_id}/download")

    assert response.status_code == 409
    assert "belum selesai" in response.json()["detail"]


def test_unknown_job_returns_404(client):
    assert client.get("/documents/jobs/tidak-ada").status_code == 404
    assert client.get("/documents/jobs/tidak-ada/download").status_code == 404


def test_bad_document_type_rejected_synchronously_without_creating_a_job(client):
    """Validasi murah tetap sinkron: request salah bentuk harus ditolak SEKARANG,
    bukan jadi job yang gagal tiga menit kemudian."""
    response = client.post(
        "/documents/generate", json={"document_type": "INVALID", "repositories": []}
    )

    assert response.status_code == 422
    assert "job_id" not in response.json()


def test_cors_exposes_content_disposition():
    """Tanpa `expose_headers`, browser MENYEMBUNYIKAN Content-Disposition dari
    JavaScript — diam-diam, tanpa error. Akibatnya frontend jatuh ke nama
    cadangan dan SETIAP pengguna mengunduh "dokumen.docx" alih-alih
    "Solution_Design_Document.docx".

    `allow_headers=["*"]` TIDAK menutupi ini: itu untuk header REQUEST.

    Bug ini ketahuan di gladi bersih demo — dari nama file yang terunduh, bukan
    dari test. Test tidak bisa menangkapnya sebelumnya karena TestClient tidak
    menegakkan CORS sama sekali: dia bukan browser. Jadi yang diperiksa di sini
    KONFIGURASINYA, dan batas itu harus jujur diakui — hanya browser sungguhan
    yang membuktikan filenya terunduh dengan nama benar."""
    from app.main import app

    cors = next(
        (m for m in app.user_middleware if "CORS" in str(m.cls)),
        None,
    )

    assert cors is not None, "CORS middleware hilang"
    exposed = cors.kwargs.get("expose_headers", [])
    assert "Content-Disposition" in exposed


# --- Progress per tahap -------------------------------------------------------
#
# Generation ~2-3 menit, dan 72%-nya ada di panggilan LLM yang tidak bisa
# dipercepat tanpa menukar kualitas dokumen — satu-satunya nilai produk ini.
# Jadi yang diperbaiki bukan durasinya, tapi keterbacaannya: `status` cuma punya
# empat nilai dan tidak bisa membedakan "sedang mengunduh repo" dari "sedang
# menunggu AI dua menit".


def test_progress_reports_each_stage_in_order(mock_plantuml_ok):
    """Urutan tahapnya = urutan pipeline. Kalau ada yang menyisipkan tahap baru
    di tempat yang salah, pengguna melihat kebohongan tentang apa yang terjadi."""
    from app.api.schemas_document import GenerateDocumentRequest

    content = _load_fixture("document_content_sdd.json")
    body = GenerateDocumentRequest(document_type="SDD", repositories=[])
    seen = []

    with patch.object(
        routes_document._llm_service, "generate_document_content", return_value=content
    ):
        routes_document._generate_document("SDD", body, on_progress=seen.append)

    assert len(seen) == 4
    assert "Mengunduh" in seen[0]
    assert seen[1].startswith("Membaca kode:")
    assert "AI" in seen[2]
    assert "diagram" in seen[3]


def test_progress_after_parse_carries_real_numbers_from_the_repo():
    """Kalimat ini muncul di detik ke-5 dan jadi bukti pertama bahwa sistem
    benar-benar membaca kode PENGGUNA — bukan mengarang. Angkanya harus nyata,
    bukan teks tetap."""
    from app.api.routes_document import _describe_parsed

    ctx = {
        "project_name": "x",
        "repositories": [
            {
                "repo_tag": "Backend",
                "files": [
                    {"api_endpoints": [1, 2, 3], "classes": [1]},
                    {"api_endpoints": [], "classes": [1, 2]},
                ],
            }
        ],
    }

    assert _describe_parsed(ctx) == "Membaca kode: 2 file, 3 endpoint, 3 class."


def test_progress_reaches_the_client_through_job_status(client):
    """Penjaga rantai: progress harus benar-benar sampai ke payload GET, bukan
    cuma tersimpan di DB."""
    job_id = job_store.create_job(document_type="SDD", project_name="x")
    job_store.mark_running(job_id)
    job_store.set_progress(job_id, "Membaca kode: 22 file, 38 endpoint, 16 class.")

    payload = client.get(f"/documents/jobs/{job_id}").json()

    assert payload["progress"] == "Membaca kode: 22 file, 38 endpoint, 16 class."


def test_init_db_adds_progress_column_to_an_existing_old_database(tmp_path, monkeypatch):
    """DB yang dibuat versi SEBELUM kolom progress ada harus ikut ditambal.

    `CREATE TABLE IF NOT EXISTS` tidak menyentuh tabel yang sudah ada, jadi tanpa
    migrasi ini DB lama meledak dengan "no such column: progress" — dan HANYA di
    mesin yang sudah pernah menjalankan versi sebelumnya, tidak pernah di test
    yang selalu mulai dari DB kosong. Bug yang tak akan terlihat di CI."""
    import sqlite3

    db = tmp_path / "lama.db"
    monkeypatch.setattr(job_store, "DB_PATH", db)
    # Skema versi lama: persis seperti sebelumnya, tanpa kolom progress.
    conn = sqlite3.connect(db)
    conn.execute(
        "CREATE TABLE jobs (id TEXT PRIMARY KEY, status TEXT NOT NULL,"
        " document_type TEXT NOT NULL, project_name TEXT, docx_path TEXT,"
        " error TEXT, error_status INTEGER, created_at TEXT NOT NULL,"
        " updated_at TEXT NOT NULL)"
    )
    conn.commit()
    conn.close()

    job_store.init_db()

    job_id = job_store.create_job(document_type="SDD", project_name="x")
    job_store.set_progress(job_id, "tahap satu")
    assert job_store.get_job(job_id)["progress"] == "tahap satu"


def test_init_db_is_safe_to_run_twice(tmp_path, monkeypatch):
    """Dipanggil tiap startup, jadi migrasi yang jalan dua kali harus tidak
    apa-apa — ALTER TABLE yang diulang akan error kalau tidak dijaga."""
    monkeypatch.setattr(job_store, "DB_PATH", tmp_path / "dua.db")

    job_store.init_db()
    job_store.init_db()

    job_id = job_store.create_job(document_type="SDD", project_name="x")
    assert job_store.get_job(job_id)["progress"] is None
