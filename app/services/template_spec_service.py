"""V2 — `TemplateSpec`: ukur docx template user jadi kontrak JSON terstruktur.

Bagian pertama dari pipeline "upload template jadi template terdaftar" (roadmap
tahap c). DETERMINISTIK, tanpa LLM — sama filosofinya dengan `parser_service`
(Contract A): ukur input, hasilkan struktur, jangan menebak.

Spec ini dikonsumsi dua langkah berikutnya:
  1. sintesis `reference.docx` (properti visual: font, heading, warna header),
  2. pemetaan bab → Contract B (outline heading berurut).

Kelayakannya sudah dibuktikan end-to-end pada template multi-sumber
(`scripts/validation/V2_TEMPLATE_MULTISOURCE.md`); modul ini mengangkat
pengukuran itu dari riset scratchpad jadi kode ber-tes.

PELAJARAN yang dikodekan di sini:
- **Font "efektif" wajib di-resolve, bukan dibaca satu sumber.** `docDefaults`
  bisa menunjuk tema sementara style `Normal` meng-override ke font lain — yang
  DIRENDER adalah hasil resolusi `Normal`→ascii/tema. Membaca `docDefaults` saja
  bisa keliru (mis. template yang docDefault-nya tema-minor tapi Normal-nya
  Verdana). Ini kelas bug yang sama dengan "font tema Aptos".
- **`caps` heading itu diukur, bukan diasumsikan** — sebagian template memaksa
  huruf besar lewat `w:caps`, sebagian tidak.
- **Judul bab tidak selalu berupa style `Heading N`.** Template buatan manusia
  menaruhnya di tempat yang tak terduga; kalau ekstraksi cuma percaya style, bab
  itu lenyap DIAM-DIAM (dokumen keluar hampa tanpa ada yang tahu). Lihat
  `_extract_outline` untuk sinyal apa saja yang dipakai dan mana yang sengaja
  TIDAK — semuanya diputuskan dari pengukuran ke template nyata, bukan tebakan.
"""
from __future__ import annotations

import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

# Style yang dicari pipeline Pandoc kita — dipakai menakar seberapa "kaya" style
# template user (V0: template nyata sering cuma punya 4 dari 14; itu sebabnya
# swap reference-doc gagal & sintesis jadi satu-satunya jalan).
PANDOC_STYLES = (
    "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4",
    "Body Text", "Compact", "First Paragraph", "Image Caption",
    "Table Caption", "Table", "TOC Heading", "toc 1", "table of figures",
)

_GENERIC_TABLE_STYLES = {"Normal Table", "Table Normal", "None"}


def _theme_fonts(path: Path) -> tuple[str | None, str | None]:
    """(font tema major, minor) dari `word/theme/theme1.xml`. python-docx tidak
    mengekspos part tema, jadi dibaca langsung dari paket zip."""
    try:
        with zipfile.ZipFile(path) as package:
            xml = package.read("word/theme/theme1.xml").decode("utf-8", "replace")
    except (KeyError, zipfile.BadZipFile, OSError):
        return (None, None)
    major = re.search(r"<a:majorFont>.*?<a:latin typeface=\"([^\"]*)\"", xml, re.S)
    minor = re.search(r"<a:minorFont>.*?<a:latin typeface=\"([^\"]*)\"", xml, re.S)
    return (major.group(1) if major else None, minor.group(1) if minor else None)


def _rfonts_value(rfonts, theme_major: str | None, theme_minor: str | None) -> str | None:
    """Font EFEKTIF dari satu elemen rFonts: `ascii` eksplisit menang; kalau
    menunjuk tema (`asciiTheme`), resolve ke tema major/minor."""
    if rfonts is None:
        return None
    ascii_ = rfonts.get(qn("w:ascii"))
    if ascii_:
        return ascii_
    theme = rfonts.get(qn("w:asciiTheme"))
    if theme:
        if "major" in theme:
            return theme_major
        if "minor" in theme:
            return theme_minor
    return None


def _docdefault_size(doc) -> float | None:
    rpr = doc.styles.element.find(
        qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    if rpr is None:
        return None
    sz = rpr.find(qn("w:sz"))
    return int(sz.get(qn("w:val"))) / 2 if sz is not None else None


def _style_element(doc, name: str):
    for style in doc.styles:
        if style.name == name:
            return style.element
    return None


def _style_rfonts(doc, name: str):
    el = _style_element(doc, name)
    if el is None:
        return None
    rpr = el.find(qn("w:rPr"))
    return rpr.find(qn("w:rFonts")) if rpr is not None else None


def _effective_fonts(doc, theme_major, theme_minor) -> tuple[str | None, str | None]:
    """(font body, font heading) yang benar-benar dirender. Body dari style
    `Normal` (fallback `docDefaults`); heading dari style `Heading 1` (fallback
    tema major). Beda body vs heading itu nyata dan harus ditangkap."""
    body = _rfonts_value(_style_rfonts(doc, "Normal"), theme_major, theme_minor)
    if not body:
        dd = doc.styles.element.find(
            qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
        dd_rfonts = dd.find(qn("w:rFonts")) if dd is not None else None
        body = _rfonts_value(dd_rfonts, theme_major, theme_minor) or theme_minor
    head = _rfonts_value(_style_rfonts(doc, "Heading 1"), theme_major, theme_minor)
    return (body, head or theme_major)


def _style_format(doc, name: str) -> dict | None:
    """Format satu style paragraf: font, ukuran (poin), bold, caps, perataan."""
    el = _style_element(doc, name)
    if el is None:
        return None
    rpr = el.find(qn("w:rPr"))
    ppr = el.find(qn("w:pPr"))
    sz = rpr.find(qn("w:sz")) if rpr is not None else None
    bold = rpr.find(qn("w:b")) if rpr is not None else None
    caps = rpr.find(qn("w:caps")) if rpr is not None else None
    rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
    jc = ppr.find(qn("w:jc")) if ppr is not None else None
    font = None
    if rfonts is not None:
        font = rfonts.get(qn("w:ascii")) or rfonts.get(qn("w:asciiTheme"))
    return {
        "font": font,
        "size_pt": int(sz.get(qn("w:val"))) / 2 if sz is not None else None,
        "bold": bold is not None and bold.get(qn("w:val")) not in ("0", "false"),
        "caps": caps is not None and caps.get(qn("w:val")) not in ("0", "false"),
        "align": jc.get(qn("w:val")) if jc is not None else None,
    }


# --- Penemuan judul bab -------------------------------------------------------
#
# Word menulis tiap shape DUA KALI demi kompatibilitas: `mc:Choice` (DrawingML
# modern) dan `mc:Fallback` (VML lama) berisi teks yang SAMA persis. Tanpa
# melewati yang Fallback, tiap judul di text box terbaca ganda (terukur pada
# `Template SDD`: 7 bab jadi 14 entri).
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"

# Judul bertingkat yang diketik manual: "1 Executive Summary", "3.1 Vision and
# Scope". Tiga penjaga, semuanya dari kasus nyata:
#
# - **Pemisah nomor-judul boleh HILANG** — Word memecah nomor dan judul ke run
#   berbeda, jadi teks gabungannya rapat ("1Executive Summary"). Tapi menerima
#   kerapatan begitu saja membuat "3D Modeling" terpotong jadi "D Modeling",
#   maka tanpa pemisah nomor harus diikuti pola KATA (huruf besar lalu kecil):
#   "1Ex…" lolos, "3D " tidak.
# - **Wajib berakhir HURUF** supaya tanggal tak lolos: "18 March 2021" pernah
#   tertangkap sebagai "judul" saat pola ini diuji ke sel tabel Veracity.
_HEADING_NUMBERED = re.compile(
    r"^(\d+(?:\.\d+)*)(?:[.)]\s*|\s+|(?=[A-Z][a-z]))([A-Za-z][^\n]{2,80}[A-Za-z])$")


def _strip_leading_number(text: str) -> str:
    """Buang penomoran yang diketik manual dari teks judul.

    Berlaku untuk heading ber-STYLE juga, bukan cuma yang dikenali lewat pola:
    tingkatnya sudah diwakili `level` dan template hasil generate menomori
    sendiri lewat tingkat heading, jadi membiarkannya menghasilkan "1Contents"
    di dokumen jadi — atau penomoran ganda ("1. 1Contents") begitu Word
    menomori otomatis.
    """
    match = _HEADING_NUMBERED.match(text)
    return match.group(2).strip() if match else text


def _element_text(element) -> str:
    return re.sub(r"\s+", " ",
                  "".join(t.text or "" for t in element.iter(qn("w:t")))).strip()


def _has_tab_char(element) -> bool:
    """Ada karakter TAB di dalam teks paragraf. Sengaja mencari `w:tab` di dalam
    `w:r` saja — `w:tab` di bawah `w:pPr` itu definisi tab STOP, bukan isi."""
    return any(run.find(qn("w:tab")) is not None
               for run in element.iter(qn("w:r")))


def _paragraph_style_name(element, style_names: dict[str, str]) -> str:
    p_pr = element.find(qn("w:pPr"))
    p_style = p_pr.find(qn("w:pStyle")) if p_pr is not None else None
    if p_style is None:
        return ""
    return style_names.get(p_style.get(qn("w:val")), "")


def _text_box_paragraphs(element):
    """Paragraf di dalam text box yang di-anchor pada `element` (tanpa kembarannya).

    `doc.paragraphs` maupun `body.iterchildren(w:p)` sama-sama BUTA terhadap ini:
    isi text box hidup di `w:txbxContent`, beberapa lapis di dalam run. Terukur
    pada `Template SDD`: 7 bab utamanya (Executive Summary ... Infrastructure)
    seluruhnya di sana, dan karena itu tak pernah masuk outline.
    """
    for box in element.iter(qn("w:txbxContent")):
        if any(a.tag == _MC_FALLBACK for a in box.iterancestors()):
            continue
        yield from box.iter(qn("w:p"))


def _paragraph_heading(element, style_names: dict[str, str], *, loose: bool):
    """`(level, teks, sinyal, nama_style)` kalau paragraf ini judul bab; None kalau bukan.

    Sinyal diurut dari yang PALING KUAT: style `Title`/`Heading N` adalah
    deklarasi eksplisit penulis, sementara pola penomoran cuma tebakan terdidik —
    jadi yang kedua hanya dipakai kalau `loose`.
    """
    name = _paragraph_style_name(element, style_names)
    text = _element_text(element)
    if name == "Title":
        return (0, _strip_leading_number(text), "style", name)
    match = re.fullmatch(r"Heading (\d)", name)
    if match is not None:
        # Tingkat dari STYLE (deklarasi penulis), teks tetap dibersihkan dari
        # nomor yang diketik manual — dua hal terpisah.
        return (int(match.group(1)), _strip_leading_number(text), "style", name)
    if loose:
        numbered = _HEADING_NUMBERED.match(text)
        if numbered is not None:
            return (numbered.group(1).count(".") + 1, numbered.group(2).strip(),
                    "numbering", name)
    return None


def _extract_outline(doc, body, orientations: list[dict]) -> dict:
    """Outline berurut + sinyal dekay, dari SELURUH wadah yang terbukti dipakai.

    Tiap entri membawa ORIENTASI section tempatnya berada. Tanpa itu `orientations`
    tak bisa dipakai apa-apa: dia daftar per-SECTION, sementara template hasil
    generate disusun per-HEADING — tak ada yang menghubungkan "section ke-2
    landscape" dengan "bab mana yang mulai di sana". Korelasinya dibuat di sini
    dengan menelusuri body secara URUT: properti sebuah section disimpan di
    `sectPr` yang MENGAKHIRInya, jadi tiap paragraf ber-`sectPr` menutup section
    berjalan dan yang berikutnya masuk section sesudahnya.

    Yang sengaja TIDAK dipindai — diukur pada 4 template nyata (`Template SDD`,
    `04. Dokumen UAT`, `Solution_Design_Document`, `Veracity`), bukan ditebak:
    - **sel tabel**: nol judul asli di keempatnya, dan satu-satunya yang mirip
      judul justru false positive ("18 March 2021"). Memindainya cuma derau.
    - **`w:outlineLvl`**: nol pemakaian di keempatnya. Sinyal sah menurut spec
      Word, tapi tak ada satu pun template kita yang bisa membuktikannya jalan —
      dan mengirim jalur yang tak terverifikasi itu justru yang kita hindari.
    """
    style_names = {s.style_id: s.name for s in doc.styles if s.style_id}
    candidates: list[dict] = []
    heading_usage = Counter()
    empty_headings = 0
    tab_in_headings = False
    section_index = 0

    def consider(element, *, where: str) -> None:
        nonlocal empty_headings, tab_in_headings
        found = _paragraph_heading(element, style_names, loose=True)
        if found is None:
            return
        level, text, signal, style_name = found
        if signal == "style":
            heading_usage[style_name] += 1
        if not text:
            empty_headings += 1
        if _has_tab_char(element):
            tab_in_headings = True
        orient = (orientations[section_index]["orient"]
                  if section_index < len(orientations) else "portrait")
        candidates.append({"level": level, "text": text[:90], "empty": not text,
                           "orient": orient, "signal": signal, "where": where})

    for element in body.iterchildren(qn("w:p")):
        consider(element, where="body")
        for nested in _text_box_paragraphs(element):
            consider(nested, where="text_box")
        p_pr = element.find(qn("w:pPr"))
        if p_pr is not None and p_pr.find(qn("w:sectPr")) is not None:
            section_index += 1

    # Pakai sinyal TERKUAT yang tersedia, jangan campur aduk. Kalau template punya
    # style heading, itu deklarasi penulis dan pola penomoran di BODY tak boleh
    # ikut campur — daftar bernomor yang diketik manual akan terseret jadi "bab".
    # Text box tetap ikut: dia buta terhadap style secara konstruksi, jadi di situ
    # penomoran bukan sinyal kedua melainkan satu-satunya.
    if any(c["signal"] == "style" for c in candidates):
        outline = [c for c in candidates
                   if c["signal"] == "style" or c["where"] == "text_box"]
    else:
        # Tak ada satu pun style heading (template yang diformat manual seluruhnya).
        # Menawarkan tebakan yang bisa dikoreksi pengguna lebih berguna daripada
        # mengembalikan outline kosong dan diam.
        outline = candidates

    return {"outline": outline,
            "heading_usage": dict(heading_usage),
            "heading_sources": dict(Counter(c["signal"] for c in outline)),
            "empty_headings": empty_headings,
            "tab_in_headings": tab_in_headings}


_UAT_MARKERS = re.compile(r"acceptance|testing|uat|pengujian|test script|defect")
_SDD_MARKERS = re.compile(r"design|architecture|deskripsi|requirement")


def _guess_kind(heading_text: str) -> str:
    """Jenis dokumen dari kosakata judul bab — yang MENANG BANYAK, bukan yang
    lebih dulu cocok.

    Versi pertama memeriksa UAT lebih dulu lalu `return` pada kecocokan PERTAMA,
    jadi satu kata "testing" di sebuah template Solution Design mengalahkan lima
    kata SDD. Bukan cacat kosmetik: jenis dokumen menentukan seluruh kosakata
    pemetaan, dan salah tebak memotong isi dokumen berlipat. Terukur pada
    `Solution_Design_Document` — dipeta sebagai UAT cuma 1 dari 17 bab terisi,
    sebagai SDD 7 dari 17.
    """
    uat = len(_UAT_MARKERS.findall(heading_text))
    sdd = len(_SDD_MARKERS.findall(heading_text))
    if uat > sdd:
        return "UAT"
    if sdd > uat:
        return "SDD"
    return "unknown"


def _guess_language(heading_text: str) -> str:
    if re.search(r"\b(dokumen|pengujian|pendahuluan|deskripsi|prosedur|"
                 r"pengendalian|persetujuan)\b", heading_text):
        return "id"
    return "en"


def build_template_spec(docx_path: str | Path) -> dict:
    """Ukur sebuah `.docx` template → `TemplateSpec` (dict). Deterministik, $0.

    `.doc` biner (format Word lama) TIDAK didukung python-docx — konversi ke
    `.docx` dulu (mis. Word COM / LibreOffice) sebelum memanggil ini.
    """
    path = Path(docx_path)
    doc = Document(str(path))
    body = doc.element.body

    theme_major, theme_minor = _theme_fonts(path)
    eff_body, eff_head = _effective_fonts(doc, theme_major, theme_minor)

    headings_fmt = {}
    for level in range(1, 5):
        fmt = _style_format(doc, f"Heading {level}")
        if fmt:
            headings_fmt[f"Heading {level}"] = fmt

    # Tabel: warna fill header baris-1 yang dominan + apakah pakai named style.
    header_fills = Counter()
    table_styles = Counter()
    for table in doc.tables:
        table_styles[table.style.name if table.style else "None"] += 1
        if not table.rows:
            continue
        for cell in table.rows[0].cells:
            shd = cell._tc.find(qn("w:tcPr") + "/" + qn("w:shd"))
            if shd is None:
                continue
            fill = shd.get(qn("w:fill"))
            if fill and fill not in ("auto", "FFFFFF"):
                header_fills[fill] += 1
    dominant_fill = header_fills.most_common(1)[0][0] if header_fills else None
    named = sum(v for k, v in table_styles.items() if k not in _GENERIC_TABLE_STYLES)
    generic = sum(v for k, v in table_styles.items() if k in _GENERIC_TABLE_STYLES)

    # Orientasi per-section.
    orientations = []
    for section in doc.sections:
        orient = ("landscape" if section.orientation == WD_ORIENT.LANDSCAPE
                  else "portrait")
        width = round(section.page_width.inches, 1) if section.page_width else None
        orientations.append({"orient": orient, "width_in": width})

    found = _extract_outline(doc, body, orientations)
    outline = found["outline"]

    n_toc = (sum(1 for it in body.findall(".//" + qn("w:instrText"))
                 if it.text and "TOC" in it.text)
             + sum(1 for fs in body.findall(".//" + qn("w:fldSimple"))
                   if "TOC" in (fs.get(qn("w:instr")) or "")))
    n_images = (len(body.findall(".//" + qn("w:drawing")))
                + len(body.findall(".//" + qn("w:pict"))))

    all_heading_text = " ".join(o["text"].lower() for o in outline)
    present_styles = {s.name for s in doc.styles}

    return {
        "source_name": path.stem,
        "doc_kind_guess": _guess_kind(all_heading_text),
        "language_guess": _guess_language(all_heading_text),
        "visual": {
            "effective_body_font": eff_body,
            "effective_heading_font": eff_head,
            "base_size_pt": _docdefault_size(doc) or 11,
            "theme_major_font": theme_major,
            "theme_minor_font": theme_minor,
            "title": _style_format(doc, "Title"),
            "headings": headings_fmt,
            "table_header_fill": dominant_fill,
            "table_header_fill_all": dict(header_fills),
            "table_uses_named_style": named > generic,
            "table_styles": dict(table_styles),
            "orientations": orientations,
            "has_landscape": any(o["orient"] == "landscape" for o in orientations),
        },
        "structure": {
            "n_paragraphs": len(doc.paragraphs),
            "n_tables": len(doc.tables),
            "n_images": n_images,
            "n_sections": len(doc.sections),
            "n_toc_fields": n_toc,
            "heading_usage": found["heading_usage"],
            "heading_sources": found["heading_sources"],
            "outline": outline,
        },
        "styles_present": [s for s in PANDOC_STYLES if s in present_styles],
        "decay": {"empty_headings": found["empty_headings"],
                  "tab_in_headings": found["tab_in_headings"]},
    }
