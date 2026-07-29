"""Tes endpoint /templates (V2 upload template). Docx sintetis (python-docx),
store di-redirect ke tmp — tanpa LLM, tanpa docx vendor. build_reference memakai
pandoc asli (seperti test compiler lain)."""
import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services import auth_service, compiler_service
from app.services.auth_service import Principal

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(autouse=True)
def isolated_store(tmp_path, monkeypatch):
    """Jangan tulis template ke data/ asli — redirect store ke tmp pytest."""
    monkeypatch.setattr(compiler_service, "TEMPLATES_STORE", tmp_path / "templates")


@pytest.fixture
def client():
    return TestClient(app)


def _docx_bytes(headings):
    doc = Document()
    doc.add_paragraph("Judul Dokumen", style="Title")
    for h in headings:
        doc.add_paragraph(h, style="Heading 1")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_compiles_registers_and_is_usable_in_generate(client):
    data = _docx_bytes(["Deskripsi Aplikasi", "Use Case", "Lampiran"])
    resp = client.post(
        "/templates",
        files={"file": ("vendor_sdd.docx", data, DOCX_MIME)},
        data={"name": "Vendor SDD", "doc_types": "SDD"},
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    tid = body["manifest"]["template_id"]
    assert tid == "vendor-sdd"
    assert body["manifest"]["doc_types"] == {"SDD": "SDD.md"}

    # rencana peta ikut di respons (untuk ditinjau)
    bindings = {m["text"]: m["binding"] for m in body["mappings"]["SDD"]}
    assert bindings["Deskripsi Aplikasi"] == "app_description"
    assert bindings["Use Case"] == "use_cases"
    assert bindings["Lampiran"] == "manual"          # tak dikenali → placeholder jujur

    # terdaftar, muncul di daftar (dengan doc_types), DAN lolos validasi generate
    listed = client.get("/templates").json()["templates"]
    entry = next(t for t in listed if t["id"] == tid)
    assert entry["doc_types"] == ["SDD"]
    assert entry["source"] == "compiled"
    compiler_service.validate_template(tid, "SDD")   # tak melempar = bisa dipakai generate

    detail = client.get(f"/templates/{tid}")
    assert detail.status_code == 200
    assert detail.json()["manifest"]["template_id"] == tid


def test_upload_with_llm_mapping_opt_in(client, monkeypatch):
    """use_llm_mapping=true meng-opt-in pemeta LLM (di-mock di boundary, $0). Bab
    asing yang heuristik tandai 'manual' kini terpetakan lewat jawaban LLM."""
    import app.services.llm_mapping_service as lm
    monkeypatch.setattr(
        lm, "_request_bindings",
        lambda outline, doc_type, usage_sink=None: {
            0: "skip", 1: "app_description", 2: "feature_requirements"},
    )
    data = _docx_bytes(["Vision Statement", "Business Capabilities"])
    resp = client.post(
        "/templates",
        files={"file": ("dynamics.docx", data, DOCX_MIME)},
        data={"doc_types": "SDD", "use_llm_mapping": "true"},
    )
    assert resp.status_code == 201, resp.text
    bindings = {m["text"]: m["binding"] for m in resp.json()["mappings"]["SDD"]}
    assert bindings["Vision Statement"] == "app_description"
    assert bindings["Business Capabilities"] == "feature_requirements"


def test_upload_rejects_doc_and_non_docx(client):
    doc = client.post("/templates", files={"file": ("lama.doc", b"\xd0\xcf", "application/msword")})
    assert doc.status_code == 422
    assert ".docx" in doc.json()["detail"]

    txt = client.post("/templates", files={"file": ("catatan.txt", b"halo", "text/plain")})
    assert txt.status_code == 422


def test_upload_rejects_garbage_and_empty(client):
    garbage = client.post(
        "/templates", files={"file": ("palsu.docx", b"ini bukan docx sama sekali", DOCX_MIME)})
    assert garbage.status_code == 422          # PackageNotFoundError → 422, bukan 500

    empty = client.post("/templates", files={"file": ("kosong.docx", b"", DOCX_MIME)})
    assert empty.status_code == 422


def test_list_templates_includes_builtins_with_doc_types(client):
    templates = client.get("/templates").json()["templates"]
    by_id = {t["id"]: t for t in templates}
    assert by_id["default"]["doc_types"] == ["SDD", "UAT"]
    assert by_id["default"]["source"] == "builtin"
    # premco kini menyediakan SDD DAN UAT — frontend membaca ini, bukan hardcode
    assert by_id["premco"]["doc_types"] == ["SDD", "UAT"]


def test_get_unknown_template_is_404(client):
    assert client.get("/templates/tidak-ada").status_code == 404


def _upload(client, headings, name="Vendor", headers=None):
    resp = client.post(
        "/templates",
        files={"file": (f"{name}.docx", _docx_bytes(headings), DOCX_MIME)},
        data={"name": name, "doc_types": "SDD"},
        headers=headers,
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def test_tinjauan_manusia_menyimpan_peta_dan_regenerate_template(client):
    """Langkah terakhir yang tak bisa diotomatiskan.

    Ekstraksi menemukan babnya dan pemeta menebak isinya, tapi cuma pemberi
    template yang tahu bab bernama asing itu sebetulnya deskripsi aplikasi."""
    body = _upload(client, ["Bab Asing Yang Tak Dikenali", "Lampiran"])
    template_id = body["manifest"]["template_id"]
    plan = body["mappings"]["SDD"]
    # Heuristik menyerah pada nama asing -> placeholder jujur, bukan tebakan.
    assert all(p["binding"] != "app_description" for p in plan)

    bindings = [p["binding"] for p in plan]
    bindings[[p["text"] for p in plan].index("Bab Asing Yang Tak Dikenali")] = "app_description"
    resp = client.put(f"/templates/{template_id}/mappings/SDD",
                      json={"bindings": bindings})

    assert resp.status_code == 200, resp.text
    updated = resp.json()["mappings"]["SDD"]
    assert updated[[p["text"] for p in plan].index("Bab Asing Yang Tak Dikenali")][
        "binding"] == "app_description"
    # Struktur bab tak boleh ikut bergeser: itu hasil PENGUKURAN dokumen sumber.
    assert [p["text"] for p in updated] == [p["text"] for p in plan]
    assert [p["level"] for p in updated] == [p["level"] for p in plan]
    # Template Jinja ikut di-generate ulang — kalau tidak, suntingan cuma tersimpan
    # di JSON dan dokumennya keluar sama saja seperti sebelum ditinjau.
    rendered = (compiler_service.TEMPLATES_STORE / template_id / "SDD.md").read_text(
        encoding="utf-8")
    assert "app_description" in rendered


def test_tinjauan_menolak_peta_yang_merusak(client):
    body = _upload(client, ["Deskripsi Aplikasi", "Use Case"])
    template_id = body["manifest"]["template_id"]
    plan = body["mappings"]["SDD"]
    bindings = [p["binding"] for p in plan]

    # Jumlah tak cocok: peta harus utuh, satu binding per bab.
    assert client.put(f"/templates/{template_id}/mappings/SDD",
                      json={"bindings": bindings[:-1]}).status_code == 422
    # Binding karangan.
    assert client.put(f"/templates/{template_id}/mappings/SDD",
                      json={"bindings": ["ngawur"] * len(bindings)}).status_code == 422
    # Satu isi dipakai dua bab: bab kedua cuma akan menyalin bab pertama, jadi
    # ditolak TERANG-TERANGAN daripada di-dedup diam-diam seperti pemeta otomatis.
    duplicated = ["app_description"] * len(bindings)
    resp = client.put(f"/templates/{template_id}/mappings/SDD",
                      json={"bindings": duplicated})
    assert resp.status_code == 422
    assert "app_description" in resp.json()["detail"]

    assert client.put("/templates/entah-apa/mappings/SDD",
                      json={"bindings": ["manual"]}).status_code == 404
    assert client.put(f"/templates/{template_id}/mappings/UAT",
                      json={"bindings": bindings}).status_code == 404


# --- Isolasi template per-pengguna (seam auth) --------------------------------

def _hs256_token(monkeypatch, sub):
    """Aktifkan auth Supabase (HS256) + kembalikan header Bearer untuk `sub`.
    Sengaja disalin dari test_routes_document: tiap berkas tes berdiri sendiri."""
    import time

    import jwt

    from app.core import config
    secret = "rahasia-test-yang-cukup-panjang-tiga-puluh-dua"
    monkeypatch.setattr(config, "SUPABASE_URL", "https://proj.supabase.co")
    monkeypatch.setattr(config, "SUPABASE_JWT_SECRET", secret)
    monkeypatch.setattr(config, "SUPABASE_JWT_AUD", "authenticated")
    token = jwt.encode({"sub": sub, "aud": "authenticated",
                        "exp": int(time.time()) + 3600}, secret, algorithm="HS256")
    return {"Authorization": f"Bearer {token}"}


def test_template_upload_terisolasi_antar_pengguna(client, monkeypatch):
    """Template hasil upload A tak boleh terlihat/terpakai oleh B.

    Yang di-upload adalah dokumen internal perusahaan pemakainya — lubang
    multi-tenant terakhir sesudah job/dokumen diisolasi. Semua penolakan berupa
    404 "tidak ditemukan", BUKAN 403: kode/pesan yang berbeda akan membocorkan
    template siapa saja yang ada di server."""
    alice = _hs256_token(monkeypatch, "alice")
    tid = _upload(client, ["Deskripsi Aplikasi", "Use Case"],
                  name="Punya Alice", headers=alice)["manifest"]["template_id"]

    # Alice: terlihat di dropdown, bisa dibaca, bisa disunting.
    listed = client.get("/templates", headers=alice).json()["templates"]
    assert tid in {t["id"] for t in listed}
    assert client.get(f"/templates/{tid}", headers=alice).status_code == 200

    # Bob: template yang sama tampak TIDAK ADA.
    bob = _hs256_token(monkeypatch, "bob")
    listed_bob = client.get("/templates", headers=bob).json()["templates"]
    assert tid not in {t["id"] for t in listed_bob}
    assert client.get(f"/templates/{tid}", headers=bob).status_code == 404
    assert client.put(f"/templates/{tid}/mappings/SDD", headers=bob,
                      json={"bindings": ["manual", "manual"]}).status_code == 404

    # ...dan tak bisa dipakai generate. Kode & pesannya PERSIS sama dengan
    # template karangan, jadi Bob tak bisa menyimpulkan template ini ada.
    asing = client.post("/documents/generate",
                        json={"document_type": "SDD", "template_id": tid}, headers=bob)
    karangan = client.post("/documents/generate",
                           json={"document_type": "SDD", "template_id": "tak-pernah-ada"},
                           headers=bob)
    assert asing.status_code == 422
    assert asing.json()["detail"].replace(tid, "X") == \
        karangan.json()["detail"].replace("tak-pernah-ada", "X")
    # Nama template Alice pun tak ikut bocor lewat daftar "tersedia: ..." di pesan.
    assert tid not in karangan.json()["detail"]


def test_builtin_tetap_milik_bersama(client, monkeypatch):
    """`default`/`premco` bukan data siapa pun — itu gaya dokumen bawaan produk,
    jadi tetap terlihat & terpakai semua pengguna."""
    alice = _hs256_token(monkeypatch, "alice")
    _upload(client, ["Deskripsi Aplikasi"], name="Punya Alice", headers=alice)

    bob = _hs256_token(monkeypatch, "bob")
    ids = {t["id"] for t in client.get("/templates", headers=bob).json()["templates"]}
    assert {"default", "premco"} <= ids
    compiler_service.validate_template("premco", "SDD", caller=None)


def test_template_lama_tanpa_owner_tetap_dipakai_siapa_pun(client, monkeypatch):
    """Kompatibilitas mundur, pola sama dengan `auth_service.owns(None, ...)`:
    template yang dikompilasi SEBELUM kolom `owner` ada (atau di mode dev) tak
    punya pemilik untuk dilanggar — jadi tetap jalan, bukan hilang diam-diam."""
    # Mode dev (auth mati, fixture conftest): upload tanpa header apa pun.
    tid = _upload(client, ["Deskripsi Aplikasi"], name="Warisan")["manifest"]["template_id"]
    manifest_path = compiler_service.TEMPLATES_STORE / tid / "template.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    del manifest["owner"]                      # persis bentuk manifest lama
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    alice = _hs256_token(monkeypatch, "alice")
    assert tid in {t["id"] for t in client.get("/templates", headers=alice).json()["templates"]}
    assert client.get(f"/templates/{tid}", headers=alice).status_code == 200
    compiler_service.validate_template(tid, "SDD", caller=Principal(
        id="alice", email=None, is_anonymous=False))


def test_mode_dev_tanpa_supabase_tak_menyaring_apa_pun(client):
    """SUPABASE_URL kosong = perilaku lama persis: semua pemanggil satu Principal
    anonim, jadi template hasil upload tetap terlihat & terpakai tanpa login."""
    tid = _upload(client, ["Deskripsi Aplikasi"], name="Dev")["manifest"]["template_id"]
    assert tid in {t["id"] for t in client.get("/templates").json()["templates"]}
    assert client.get(f"/templates/{tid}").status_code == 200
    compiler_service.validate_template(tid, "SDD", caller=auth_service.ANONYMOUS)


def test_upload_template_kena_rate_limit_per_akun(client, monkeypatch):
    """Upload dibatasi juga: `use_llm_mapping=true` BERBAYAR, dan bahkan jalur
    $0-nya memakan CPU/disk untuk file sampai 50 MB. Batasnya per-akun."""
    from app.core import config
    monkeypatch.setattr(config, "RATE_LIMIT_TEMPLATE_UPLOAD_PER_WINDOW", 1)
    alice = _hs256_token(monkeypatch, "alice")

    _upload(client, ["Deskripsi Aplikasi"], name="Pertama", headers=alice)
    ditolak = client.post(
        "/templates",
        files={"file": ("kedua.docx", _docx_bytes(["Use Case"]), DOCX_MIME)},
        data={"name": "Kedua", "doc_types": "SDD"},
        headers=alice,
    )
    assert ditolak.status_code == 429
    assert int(ditolak.headers["Retry-After"]) > 0
    # Template kedua tak pernah dikompilasi — 429 berarti kerjanya tak dikerjakan.
    assert not (compiler_service.TEMPLATES_STORE / "kedua").exists()

    # Akun lain tak ikut terhukum.
    bob = _hs256_token(monkeypatch, "bob")
    _upload(client, ["Deskripsi Aplikasi"], name="Punya Bob", headers=bob)


def test_daftar_binding_tersedia_untuk_dropdown(client):
    """Path statis harus menang atas `/{template_id}` — kalau urutan deklarasinya
    terbalik, parameter menelan "bindings" dan endpoint ini balas 404 diam-diam."""
    resp = client.get("/templates/bindings")
    assert resp.status_code == 200
    bindings = resp.json()["bindings"]
    assert {"skip", "manual", "heading_only", "app_description"} <= set(bindings)
