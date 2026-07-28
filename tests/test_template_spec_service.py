"""Tes `template_spec_service` (V2 increment 1). Fixture SINTETIS dibangun
python-docx — bukan docx vendor (yang gitignore) — supaya jalan di CI mana pun."""
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.services.template_spec_service import build_template_spec


def _shade_first_row(table, fill):
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)


def _build_synthetic_docx(path):
    doc = Document()
    h1 = doc.styles["Heading 1"]
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.all_caps = True
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph("Judul Dokumen", style="Title")
    doc.add_paragraph("Bab Satu", style="Heading 1")
    doc.add_paragraph("Sub A", style="Heading 2")
    doc.add_paragraph("Bab Dua", style="Heading 1")
    doc.add_paragraph("", style="Heading 1")   # heading kosong -> dekay
    doc.add_paragraph("Paragraf body biasa untuk mengisi dokumen.")

    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "No"
    _shade_first_row(table, "A4A4A4")

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE

    doc.save(str(path))
    return path


def test_build_template_spec_menangkap_struktur_dan_visual(tmp_path):
    path = _build_synthetic_docx(tmp_path / "synthetic.docx")
    spec = build_template_spec(path)

    # --- struktur / outline berurut ---
    st = spec["structure"]
    assert st["n_tables"] == 1
    assert st["n_sections"] == 2
    outline = [(o["level"], o["text"]) for o in st["outline"]]
    assert (0, "Judul Dokumen") in outline
    assert (1, "Bab Satu") in outline
    assert (2, "Sub A") in outline
    assert st["heading_usage"]["Heading 1"] == 3   # Bab Satu, Bab Dua, kosong
    assert st["heading_usage"]["Heading 2"] == 1

    # --- dekay: satu heading kosong ---
    assert spec["decay"]["empty_headings"] == 1

    # --- visual: heading terukur (bukan diasumsikan) ---
    h1 = spec["visual"]["headings"]["Heading 1"]
    assert h1["size_pt"] == 16
    assert h1["bold"] is True
    assert h1["caps"] is True
    assert h1["align"] == "left"

    # --- visual: warna header tabel + orientasi landscape ---
    assert spec["visual"]["table_header_fill"] == "A4A4A4"
    assert spec["visual"]["has_landscape"] is True

    # --- font efektif ter-resolve (non-None) ---
    assert spec["visual"]["effective_body_font"]
    assert spec["visual"]["effective_heading_font"]


def test_build_template_spec_dokumen_polos_tanpa_error(tmp_path):
    """Dokumen minimal (tanpa tabel/heading khusus) tetap menghasilkan spec utuh,
    tidak melempar — jalur 'template miskin' yang V0 catat (4/14 style)."""
    doc = Document()
    doc.add_paragraph("Cuma paragraf biasa.")
    p = tmp_path / "polos.docx"
    doc.save(str(p))

    spec = build_template_spec(p)
    assert spec["structure"]["n_tables"] == 0
    assert spec["structure"]["outline"] == []
    assert spec["visual"]["table_header_fill"] is None
    assert spec["visual"]["has_landscape"] is False
    assert spec["visual"]["base_size_pt"]        # selalu ada fallback


def test_outline_entries_carry_their_section_orientation(tmp_path):
    """Bab harus tahu dirinya berada di section potret atau landscape.

    `orientations` sendiri sudah lama diukur, tapi itu daftar per-SECTION
    sementara template yang di-generate disusun per-HEADING — tanpa korelasi ini
    pengukurannya tak pernah bisa dipakai, dan tabel lebar milik template
    pengguna dipaksa muat di halaman potret."""
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches

    d = docx.Document()
    d.add_heading("Pendahuluan", 1)
    landscape = d.add_section()
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = Inches(11), Inches(8.5)
    d.add_heading("Matriks Pengujian", 1)
    portrait = d.add_section()
    portrait.orientation = WD_ORIENT.PORTRAIT
    portrait.page_width, portrait.page_height = Inches(8.5), Inches(11)
    d.add_heading("Penutup", 1)
    source = tmp_path / "tiga_section.docx"
    d.save(str(source))

    spec = build_template_spec(source)

    by_text = {o["text"]: o["orient"] for o in spec["structure"]["outline"]}
    assert by_text == {
        "Pendahuluan": "portrait",
        "Matriks Pengujian": "landscape",
        "Penutup": "portrait",
    }
    assert spec["visual"]["has_landscape"] is True


# --- Judul bab di luar style `Heading N` ---------------------------------------
#
# Template buatan manusia sering tak memakai style heading sama sekali untuk bab
# utamanya. Kasus nyata yang melahirkan tes-tes ini: sebuah template SDD vendor
# menyimpan KETUJUH bab utamanya di text box tanpa style apa pun — dan karena
# `doc.paragraphs` buta terhadap text box, dokumen hasil generate kehilangan
# ketujuhnya TANPA SATU PUN PESAN. Yang paling mahal bukan babnya hilang,
# melainkan hilangnya senyap.

_TEXT_BOX_XML = """
<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <w:r>
    <mc:AlternateContent>
      <mc:Choice Requires="wps">
        <wps:txbx><w:txbxContent>
          <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
        </w:txbxContent></wps:txbx>
      </mc:Choice>
      <mc:Fallback>
        <w:pict><v:textbox><w:txbxContent>
          <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
        </w:txbxContent></v:textbox></w:pict>
      </mc:Fallback>
    </mc:AlternateContent>
  </w:r>
</w:p>
"""


def _add_text_box(doc, text):
    """Sisipkan paragraf ber-text box, DUA salinan seperti Word menulisnya.

    Word memancarkan tiap shape dua kali (`mc:Choice` DrawingML + `mc:Fallback`
    VML lama) dengan teks identik. Fixture ini sengaja meniru itu — kalau tidak,
    tes akan lulus sambil membiarkan bug "tiap judul terbaca dua kali" lolos.
    """
    from docx.oxml import parse_xml

    element = parse_xml(_TEXT_BOX_XML.format(text=text))
    sect_pr = doc.element.body.find(qn("w:sectPr"))
    if sect_pr is not None:
        sect_pr.addprevious(element)
    else:
        doc.element.body.append(element)


def test_judul_di_text_box_terbaca_sekali_dengan_tingkat_dari_nomornya(tmp_path):
    doc = Document()
    doc.add_paragraph("Judul", style="Title")
    _add_text_box(doc, "1Executive Summary")      # Word merapatkan nomor & judul
    _add_text_box(doc, "2 Business Situation")
    doc.add_paragraph("Latar", style="Heading 2")
    path = tmp_path / "textbox.docx"
    doc.save(str(path))

    outline = build_template_spec(path)["structure"]["outline"]
    from_box = [o for o in outline if o["where"] == "text_box"]

    # Sekali, bukan dua kali: salinan `mc:Fallback` harus dilewati.
    assert [o["text"] for o in from_box] == ["Executive Summary", "Business Situation"]
    assert [o["level"] for o in from_box] == [1, 1]
    # Urutan dokumen tetap terjaga: bab text box mendahului sub-babnya yang ber-style,
    # jadi hierarki induk-anak di template hasil generate tidak terbalik.
    assert [o["text"] for o in outline][-3:] == [
        "Executive Summary", "Business Situation", "Latar"]


def test_tanggal_di_text_box_bukan_judul(tmp_path):
    """Pola penomoran harus menolak tanggal — "18 March 2021" pernah tertangkap
    sebagai judul saat pola ini diuji ke template nyata."""
    doc = Document()
    doc.add_paragraph("Bab", style="Heading 1")
    _add_text_box(doc, "18 March 2021")
    path = tmp_path / "tanggal.docx"
    doc.save(str(path))

    outline = build_template_spec(path)["structure"]["outline"]
    assert [o["text"] for o in outline] == ["Bab"]


def test_daftar_bernomor_di_body_tidak_dianggap_bab(tmp_path):
    """Kalau template PUNYA style heading, itu deklarasi eksplisit penulis dan
    pola penomoran tak boleh ikut campur di body — daftar bernomor yang diketik
    manual akan terseret jadi "bab" dan mengacaukan seluruh peta."""
    doc = Document()
    doc.add_paragraph("Prosedur", style="Heading 1")
    doc.add_paragraph("1 Buka aplikasi lalu masuk")
    doc.add_paragraph("2 Pilih menu Laporan")
    path = tmp_path / "daftar.docx"
    doc.save(str(path))

    structure = build_template_spec(path)["structure"]
    assert [o["text"] for o in structure["outline"]] == ["Prosedur"]
    assert structure["heading_sources"] == {"style": 1}


def test_template_tanpa_style_heading_jatuh_ke_pola_penomoran(tmp_path):
    """Template yang diformat manual seluruhnya: menawarkan tebakan yang bisa
    dikoreksi pengguna lebih berguna daripada mengembalikan outline kosong."""
    doc = Document()
    doc.add_paragraph("1 Pendahuluan")
    doc.add_paragraph("1.1 Latar Belakang")
    doc.add_paragraph("Paragraf isi yang panjangnya wajar dan bukan judul apa pun.")
    doc.add_paragraph("2 Ruang Lingkup")
    path = tmp_path / "polos.docx"
    doc.save(str(path))

    structure = build_template_spec(path)["structure"]
    assert [(o["level"], o["text"]) for o in structure["outline"]] == [
        (1, "Pendahuluan"), (2, "Latar Belakang"), (1, "Ruang Lingkup")]
    assert structure["heading_sources"] == {"numbering": 3}


def test_jenis_dokumen_dipilih_dari_sinyal_terkuat_bukan_yang_pertama_cocok(tmp_path):
    """Satu kata "testing" tak boleh mengalahkan lima kata Solution Design.

    Bukan cacat kosmetik: jenis dokumen menentukan seluruh kosakata pemetaan.
    Terukur pada template nyata — salah tebak memotong bab yang terisi dari 7
    jadi 1 dari 17."""
    doc = Document()
    for text in ["Solution Design Overview", "System Architecture",
                 "Functional Requirement", "Design Consideration",
                 "Testing Notes"]:
        doc.add_paragraph(text, style="Heading 1")
    path = tmp_path / "sdd_dengan_satu_kata_testing.docx"
    doc.save(str(path))

    assert build_template_spec(path)["doc_kind_guess"] == "SDD"


def test_nomor_yang_diketik_manual_dibuang_dari_judul_termasuk_yang_ber_style(tmp_path):
    """Tingkat bab datang dari style/pola; nomornya TIDAK boleh ikut jadi teks.

    Template hasil generate menomori sendiri lewat tingkat heading, jadi nomor
    yang tertinggal menghasilkan "1Contents" di dokumen jadi — atau penomoran
    ganda ("1. 1Contents") begitu Word menomori otomatis."""
    doc = Document()
    doc.add_paragraph("1Contents", style="Heading 1")      # Word merapatkan run
    doc.add_paragraph("2.1 Ruang Lingkup", style="Heading 2")
    doc.add_paragraph("3D Modeling", style="Heading 1")    # BUKAN nomor bab
    path = tmp_path / "nomor.docx"
    doc.save(str(path))

    outline = build_template_spec(path)["structure"]["outline"]

    assert [o["text"] for o in outline] == ["Contents", "Ruang Lingkup", "3D Modeling"]
    # Tingkat tetap dari STYLE, bukan dari jumlah titik pada nomornya.
    assert [o["level"] for o in outline] == [1, 2, 1]
